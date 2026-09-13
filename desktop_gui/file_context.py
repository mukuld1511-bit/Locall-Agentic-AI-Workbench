from __future__ import annotations

from pathlib import Path


def read_file(path_str: str) -> str:
    path = Path(path_str).expanduser().resolve()

    if not path.exists():
        raise FileNotFoundError(path)

    if not path.is_file():
        raise ValueError("Not a file")

    suffix = path.suffix.lower()

    if suffix in {
        ".txt",
        ".md",
        ".csv",
        ".json",
        ".xml",
        ".yaml",
        ".yml",
        ".log",
        ".py",
    }:
        return path.read_text(
            encoding="utf-8",
            errors="replace",
        )[:150000]

    if suffix == ".pdf":
        import fitz

        doc = fitz.open(str(path))

        return "\n\n".join(
            f"--- PAGE {i+1} ---\n{page.get_text()}"
            for i, page in enumerate(doc)
        )[:180000]

    if suffix == ".docx":
        import docx

        doc = docx.Document(str(path))
        parts = []

        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text)

        for table in doc.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(
                        cell.text
                        for cell in row.cells
                    )
                )

        return "\n".join(parts)[:180000]

    if suffix == ".xlsx":
        import openpyxl

        wb = openpyxl.load_workbook(
            str(path),
            read_only=True,
            data_only=True,
        )

        parts = []

        for ws in wb.worksheets:
            parts.append(
                f"--- SHEET: {ws.title} ---"
            )

            for row in ws.iter_rows(
                values_only=True
            ):
                vals = [
                    "" if x is None else str(x)
                    for x in row
                ]

                if any(vals):
                    parts.append(
                        " | ".join(vals)
                    )

        return "\n".join(parts)[:180000]

    raise ValueError(
        f"Unsupported text file type: {suffix}"
    )
