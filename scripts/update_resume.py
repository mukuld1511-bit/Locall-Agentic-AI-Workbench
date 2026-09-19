import docx
from docx.shared import Pt, Inches, RGBColor
from pathlib import Path

source_path = Path(r"C:\AI\Locall-Agentic-AI-Workbench\Mukul_Dhankhar_Resume_Updated (2).docx")
backup_path = Path(r"C:\AI\Locall-Agentic-AI-Workbench\Mukul_Dhankhar_Resume_Backup.docx")
target_path = Path(r"C:\AI\Locall-Agentic-AI-Workbench\Mukul_Dhankhar_Resume_Updated.docx")

# Backup original
if not backup_path.exists():
    import shutil
    shutil.copy2(source_path, backup_path)
    print(f"[*] Created backup at {backup_path}")

doc = docx.Document(source_path)

# Updated bullets for Sovereign Industrial AI Workbench
new_project_title = "Sovereign Industrial AI Workbench — Air-Gapped Multi-Model Operating System"
new_bullets = [
    "Architected an air-gapped, zero-cloud sovereign industrial AI workbench for mission-critical petroleum refinery operations (MRPL SIH26117), eliminating external telemetry and egress.",
    "Engineered a dynamic 4-tier model orchestration engine with automatic GPU/VRAM auto-budgeting, orchestrating a fine-tuned 500M organizer, 3B code generator, and 3B multimodal vision VLM (Qwen2.5-VL) for real-time P&ID blueprint analysis.",
    "Integrated a 3D digital twin (Three.js WebGL) with simulated combustion blower flames, rotating impellers, and high-voltage electrical plasma synchronized with live SCADA sensor telemetry and multi-lingual voice actuation (English/Hindi/Hinglish).",
    "Developed a cryptographically chained SHA-256 tamper-evident audit ledger, automated fail-safe emergency trips, and a 4-tier RBAC policy gateway enforcing strict default-deny safety interlocks across plant machinery and database operations."
]

# Update P[25] (Title)
p25 = doc.paragraphs[25]
p25.clear()
r1 = p25.add_run(new_project_title)
r1.bold = True
r1.font.size = Pt(10.5)
r2 = p25.add_run("\tSep 2026 - Present")
r2.font.size = Pt(10)

# Replace bullets 26, 27, 28, 29
for idx, bullet_text in enumerate(new_bullets):
    p = doc.paragraphs[26 + idx]
    p.clear()
    r = p.add_run(bullet_text)
    r.font.size = Pt(10)

# Update Technical Skills P[49] and P[51] to highlight Electron, Three.js, Native Linux/DGX deployment
# P[49]: AI Infrastructure & Agentic Systems
p49 = doc.paragraphs[49]
p49.clear()
r_label = p49.add_run("AI Infrastructure & Agentic Systems: ")
r_label.bold = True
r_label.font.size = Pt(10)
r_val = p49.add_run("On-Premise / Air-Gapped LLM & VLM Deployment, NVIDIA DGX / CUDA (Multi-GPU), llama.cpp, Local Model Quantization (GGUF), Model Orchestration & Routing, Multi-Agent Systems, RAG, Central Policy Engine & RBAC")
r_val.font.size = Pt(10)

# P[51]: Web, Desktop & DevOps
p51 = doc.paragraphs[51]
p51.clear()
r_label2 = p51.add_run("Web, Desktop & DevOps: ")
r_label2.bold = True
r_label2.font.size = Pt(10)
r_val2 = p51.add_run("FastAPI, React, TypeScript, Vite, Electron, Monaco Editor, TailwindCSS, Three.js (WebGL 3D) — SQLite, PostgreSQL — Linux / DGX Spark Native Deployment, Docker, Git")
r_val2.font.size = Pt(10)

# Save to target and overwrite main
doc.save(target_path)
doc.save(source_path)
print(f"[+] Successfully updated resume at {source_path} and {target_path}!")
